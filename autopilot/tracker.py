"""Local records: seen postings, the job list, applications.xlsx and Dashboard.html."""
import datetime as dt
import json
import re
from pathlib import Path

from .config import ROOT


def load(path, default):
    try:
        return json.loads(Path(path).read_text())
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def save(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=1, ensure_ascii=False))


def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()


def dedup_key(company, title):
    t = re.sub(r"\b(summer|fall|spring|winter|20\d\d|intern(ship)?)\b", " ", norm(title))
    return norm(company) + "|" + " ".join(t.split())


class Store:
    def __init__(self, cfg):
        self.out = cfg["paths"]["output_dir"]
        self.cfg = cfg
        self.jobs_path = self.out / "data" / "jobs.json"
        self.seen_path = self.out / "data" / "seen.json"
        self.jobs = load(self.jobs_path, [])
        self.seen = load(self.seen_path, {})

    def is_seen(self, c):
        keys = {v.get("key") for v in self.seen.values()}
        return c["id"] in self.seen or dedup_key(c["company"], c["title"]) in keys

    def mark(self, c, status, reason=""):
        self.seen[c["id"]] = {"key": dedup_key(c["company"], c["title"]), "status": status,
                              "reason": reason, "date": dt.date.today().isoformat()}
        save(self.seen_path, self.seen)

    def add(self, job):
        self.jobs = [j for j in self.jobs if j["id"] != job["id"]] + [job]
        save(self.jobs_path, self.jobs)
        self.mark(job, "tailored")

    # ---- statuses saved from Dashboard.html ("Save statuses" downloads autopilot-status.json)
    def import_statuses(self):
        folder = self.cfg["paths"]["status_import_dir"]
        files = sorted(folder.glob("autopilot-status*.json"), key=lambda f: f.stat().st_mtime)
        by_id, n = {j["id"]: j for j in self.jobs}, 0
        for f in files:
            for jid, v in load(f, {}).get("statuses", {}).items():
                j = by_id.get(jid)
                if j and v.get("status") and v.get("at", "") > j.get("statusUpdated", ""):
                    j["status"], j["statusUpdated"] = v["status"], v["at"]
                    n += 1
        if n:
            save(self.jobs_path, self.jobs)
        return n

    def build(self):
        imported = self.import_statuses()
        self.write_dashboard()
        xlsx = self.write_xlsx()
        print(f"Dashboard.html{' + applications.xlsx' if xlsx else ''}: {len(self.jobs)} postings"
              + (f"; imported {imported} saved statuses" if imported else ""))

    def write_dashboard(self):
        rows = []
        for j in self.jobs:
            rows.append({k: j.get(k, "") for k in ("id", "date", "company", "title", "location", "url",
                                                  "role_type", "priority", "fit", "visa_note", "notes",
                                                  "status", "statusUpdated", "pdf", "jd")})
        data = {"generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
                "subtitle": self.cfg.get("dashboard", {}).get("subtitle", ""),
                "jobs": rows}
        blob = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
        tpl = (ROOT / "templates" / "dashboard.html").read_text()
        (self.out / "Dashboard.html").write_text(tpl.replace("/*__DATA__*/null", blob))

    def write_xlsx(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            return False
        wb = Workbook()
        ws = wb.active
        ws.title = "Applications"
        cols = [("Date found", 12), ("Company", 22), ("Title", 44), ("Location", 24), ("Type", 9),
                ("Fit", 6), ("Visa note", 26), ("Status", 12), ("Posting", 12), ("Resume", 12), ("Notes", 70)]
        for i, (name, w) in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
            ws.column_dimensions[get_column_letter(i)].width = w
        for r, j in enumerate(sorted(self.jobs, key=lambda j: j["date"], reverse=True), 2):
            vals = [j["date"], j["company"], j["title"], j["location"], j["role_type"], j["fit"],
                    j.get("visa_note", ""), j.get("status", "To apply"), "Open", "Open", j["notes"]]
            for c, v in enumerate(vals, 1):
                ws.cell(row=r, column=c, value=v).alignment = Alignment(vertical="top", wrap_text=c in (3, 7, 11))
            ws.cell(row=r, column=9).hyperlink = j["url"]
            ws.cell(row=r, column=10).hyperlink = (self.out / j["pdf"]).as_uri()
            for c in (9, 10):
                ws.cell(row=r, column=c).font = Font(color="0563C1", underline="single")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(self.out / "applications.xlsx")
        return True
